"""Практика №2: Excel-книга АНР-модели с живыми формулами.

Запускать после scripts/p2_ahp.py. Результат: outputs/p2/ahp_model.xlsx
Жёлтые ячейки - суждения ЛПР (их можно менять), синие числа - исходные данные,
чёрные - формулы.
"""
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "outputs" / "p2"

FONT = "Times New Roman"
F_TEXT = Font(name=FONT, size=11)
F_BOLD = Font(name=FONT, size=11, bold=True)
F_TITLE = Font(name=FONT, size=13, bold=True)
F_INPUT = Font(name=FONT, size=11, color="0000FF")
F_LINK = Font(name=FONT, size=11, color="008000")
YELLOW = PatternFill("solid", start_color="FFFF00")
GREY = PatternFill("solid", start_color="EDEDED")
GREEN_FILL = PatternFill("solid", start_color="C6EFCE")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def cell(ws, ref, value, font=F_TEXT, fill=None, fmt=None, align=CENTER, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    c.alignment = align
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    return c


def sheet_instructions(wb):
    ws = wb.active
    ws.title = "Инструкция"
    lines = [
        ("АНР-модель выбора поставщиков оборудования для диспетчерских комплексов (практика № 2)", F_TITLE),
        ("АНР - метод анализа иерархий: критерии сравниваются попарно по шкале Саати 1-9, из матрицы получаются веса.", F_TEXT),
        ("Обозначения цветом:", F_BOLD),
        ("  жёлтая ячейка, синий текст - суждение ЛПР (можно менять, лист «Критерии», верхний треугольник матрицы);", F_TEXT),
        ("  синий текст - исходные данные (цены, сроки, оценки 1-5) и значения, рассчитанные в Python;", F_TEXT),
        ("  зелёный текст - ссылка на другой лист; чёрный текст - формула.", F_TEXT),
        ("Шкала Саати: 1 - равная важность, 3 - умеренное превосходство, 5 - существенное, 7 - значительное, 9 - абсолютное; 2, 4, 6, 8 - промежуточные; обратные дроби - если важнее второй критерий.", F_TEXT),
        ("ИС - индекс согласованности, СИ - случайный индекс Саати, ОС - отношение согласованности = ИС / СИ; суждения согласованы при ОС ≤ 0,10.", F_TEXT),
        ("Локальные приоритеты альтернатив рассчитаны в Python (scripts/p2_ahp.py) по матрицам парных сравнений альтернатив; пример такой матрицы на формулах - лист «Пример_ИБП_k1».", F_TEXT),
        ("Оценки 1-5 (качество, надёжность, санкционный риск, происхождение) - экспертные оценки аналитика по правилам из открытых сведений (sppr/ratings.py, config/quality_ratings.json).", F_TEXT),
        ("Пример изменения: если директор считает стоимость в 3 раза важнее качества, введите 3 в ячейку «k1 / k2» - веса, ОС и ранги пересчитаются.", F_TEXT),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130


def sheet_criteria(wb, criteria, judgments, py_weights):
    ws = wb.create_sheet("Критерии")
    n = len(criteria)
    ids = [c["id"] for c in criteria]
    first, last = 4, 3 + n                      # строки матрицы
    mcol = lambda j: get_column_letter(2 + j)  # колонки B..H
    gcol, wcol, awcol, ratio_col, pycol = (get_column_letter(2 + n + k) for k in range(5))
    cell(ws, "A1", "Матрица парных сравнений критериев (шкала Саати 1-9)", F_TITLE, align=LEFT, border=False)
    cell(ws, "A2", "Жёлтые ячейки - суждения ЛПР; нижний треугольник = 1 / верхний.", align=LEFT, border=False)
    cell(ws, "A3", "Критерий", F_BOLD, GREY)
    for j, cid in enumerate(ids):
        cell(ws, f"{mcol(j)}3", cid, F_BOLD, GREY)
    for ref, title in ((gcol, "Среднее геометрическое строки"), (wcol, "Вес w"), (awcol, "(A·w)"),
                       (ratio_col, "(A·w) / w"), (pycol, "Вес (Python, собств. вектор)")):
        cell(ws, f"{ref}3", title, F_BOLD, GREY)
    for i, crit in enumerate(criteria):
        r = first + i
        cell(ws, f"A{r}", f"{crit['id']} {crit['name']}", F_BOLD, align=LEFT)
        for j in range(n):
            ref = f"{mcol(j)}{r}"
            if i == j:
                cell(ws, ref, 1, fmt="0")
            elif i < j:
                c = cell(ws, ref, judgments[(i, j)], F_INPUT, YELLOW, fmt="0.00")
                c.comment = Comment(f"Суждение ЛПР: во сколько раз {ids[i]} важнее {ids[j]}", "СППР")
            else:
                cell(ws, ref, f"=1/{mcol(i)}{first + j}", fmt="0.00")
        row_range = f"{mcol(0)}{r}:{mcol(n - 1)}{r}"
        cell(ws, f"{gcol}{r}", f"=PRODUCT({row_range})^(1/{n})", fmt="0.0000")
        cell(ws, f"{wcol}{r}", f"={gcol}{r}/SUM(${gcol}${first}:${gcol}${last})", F_BOLD, fmt="0.000")
        terms = "+".join(f"{mcol(j)}{r}*${wcol}${first + j}" for j in range(n))
        cell(ws, f"{awcol}{r}", f"={terms}", fmt="0.0000")
        cell(ws, f"{ratio_col}{r}", f"={awcol}{r}/{wcol}{r}", fmt="0.0000")
        cell(ws, f"{pycol}{r}", float(py_weights[i]), F_INPUT, fmt="0.000")
    s = last + 2
    rows = [
        ("n (число критериев)", n, "0", F_INPUT),
        ("λmax = среднее (A·w)/w", f"=AVERAGE({ratio_col}{first}:{ratio_col}{last})", "0.0000", F_TEXT),
        ("ИС = (λmax - n) / (n - 1)", f"=(B{s + 1}-B{s})/(B{s}-1)", "0.0000", F_TEXT),
        ("СИ (случайный индекс Саати для n = 7)", 1.32, "0.00", F_INPUT),
        ("ОС = ИС / СИ", f"=B{s + 2}/B{s + 3}", "0.0000", F_BOLD),
        ("Вывод", f'=IF(B{s + 4}<=0.1,"суждения согласованы (ОС ≤ 0,10)","пересмотреть суждения (ОС > 0,10)")', None, F_BOLD),
    ]
    for k, (label, value, fmt, font) in enumerate(rows):
        cell(ws, f"A{s + k}", label, F_BOLD if k in (4, 5) else F_TEXT, GREY, align=LEFT)
        cell(ws, f"B{s + k}", value, font, fmt=fmt)
    ws[f"B{s + 3}"].comment = Comment("Случайный индекс Саати: n=3 0,58; 4 0,90; 5 1,12; 6 1,24; 7 1,32", "СППР")
    ws.column_dimensions["A"].width = 38
    for j in range(n):
        ws.column_dimensions[mcol(j)].width = 8
    for ref in (gcol, wcol, awcol, ratio_col, pycol):
        ws.column_dimensions[ref].width = 15
    ws.row_dimensions[3].height = 45
    return {"weight_cells": [f"Критерии!${wcol}${first + i}" for i in range(n)], "cr_cell": f"Критерии!$B${s + 4}"}


def sheet_data(wb, scored):
    ws = wb.create_sheet("Данные")
    cols = [("category", "Категория", 14), ("vendor", "Поставщик", 14), ("model", "Модель", 34),
            ("price_rub", "k1 Цена, ₽", 13), ("k2_quality", "k2 Качество, 1-5", 10),
            ("k3_reliability", "k3 Надёжность, 1-5", 11), ("lead_time_weeks_model", "k4 Срок, нед.", 9),
            ("warranty_months_model", "k5 Гарантия, мес.", 10), ("k6_sanction_risk", "k6 Санкц. риск, 1-5", 11),
            ("k7_russian_origin", "k7 Происхождение, 1-5", 12), ("price_status", "Статус цены", 18),
            ("price_url", "Источник цены", 40), ("notes", "Оговорки", 40)]
    cell(ws, "A1", "Исходные данные по альтернативам (открытые источники; оценки 1-5 - экспертные оценки аналитика)", F_TITLE, align=LEFT, border=False)
    for j, (_, title, width) in enumerate(cols, start=1):
        cell(ws, f"{get_column_letter(j)}3", title, F_BOLD, GREY)
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(scored.to_dict("records")):
        r = 4 + i
        for j, (key, _, _) in enumerate(cols, start=1):
            v = row.get(key, "")
            is_num = key in ("price_rub", "k2_quality", "k3_reliability", "lead_time_weeks_model",
                             "warranty_months_model", "k6_sanction_risk", "k7_russian_origin")
            c = cell(ws, f"{get_column_letter(j)}{r}", float(v) if is_num else str(v),
                     F_INPUT if is_num else F_TEXT, fmt="#,##0" if key == "price_rub" else ("0" if is_num else None),
                     align=CENTER if is_num else LEFT)
        for key, col in (("k3_why", 6), ("lead_time_why", 7), ("warranty_why", 8), ("k6_why", 9), ("k7_why", 10), ("k2_why", 5)):
            if row.get(key):
                ws.cell(row=r, column=col).comment = Comment(str(row[key]), "СППР")
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "D4"


def sheet_priorities(wb, local, crit_ids, weight_cells):
    ws = wb.create_sheet("Приоритеты")
    n = len(crit_ids)
    cell(ws, "A1", "Локальные приоритеты (Python, собственный вектор матриц парных сравнений альтернатив) и глобальный приоритет (формула)", F_TITLE, align=LEFT, border=False)
    cell(ws, "A2", "Глобальный приоритет = Σ w_k · p_k; веса берутся с листа «Критерии», ранг - внутри категории.", align=LEFT, border=False)
    heads = ["Категория", "Поставщик"] + [f"p({c})" for c in crit_ids] + ["Глобальный приоритет", "Ранг в категории"]
    for j, h in enumerate(heads, start=1):
        cell(ws, f"{get_column_letter(j)}3", h, F_BOLD, GREY)
    cell(ws, "A4", "Вес w →", F_BOLD, GREY, align=LEFT)
    cell(ws, "B4", "", fill=GREY)
    for k in range(n):
        cell(ws, f"{get_column_letter(3 + k)}4", f"={weight_cells[k]}", F_LINK, fmt="0.000")
    gcol, rcol = get_column_letter(3 + n), get_column_letter(4 + n)
    first = 5
    last = first + len(local) - 1
    for i, row in enumerate(local.to_dict("records")):
        r = first + i
        cell(ws, f"A{r}", row["category"], align=LEFT)
        cell(ws, f"B{r}", row["vendor"], align=LEFT)
        for k, cid in enumerate(crit_ids):
            cell(ws, f"{get_column_letter(3 + k)}{r}", float(row[cid]), F_INPUT, fmt="0.000")
        terms = "+".join(f"{get_column_letter(3 + k)}{r}*{get_column_letter(3 + k)}$4" for k in range(n))
        cell(ws, f"{gcol}{r}", f"={terms}", F_BOLD, fmt="0.000")
        cell(ws, f"{rcol}{r}", f"=COUNTIFS($A${first}:$A${last},A{r},${gcol}${first}:${gcol}${last},\">\"&{gcol}{r})+1", fmt="0")
    ws.conditional_formatting.add(f"A{first}:{rcol}{last}", FormulaRule(formula=[f"${rcol}{first}=1"], fill=GREEN_FILL))
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    for k in range(n + 2):
        ws.column_dimensions[get_column_letter(3 + k)].width = 12
    ws.row_dimensions[3].height = 35
    ws.freeze_panes = "C5"


def sheet_example(wb, scored):
    """Полная матрица сравнения ИБП по стоимости (k1, min) на формулах."""
    ws = wb.create_sheet("Пример_ИБП_k1")
    data_rows = [4 + pos for pos, cat in enumerate(scored["category"]) if cat == "ИБП"]  # строки на листе «Данные»
    ups = scored[scored["category"] == "ИБП"].reset_index(drop=True)
    m = len(ups)
    cell(ws, "A1", "Пример: матрица парных сравнений ИБП по стоимости (k1, чем дешевле - тем лучше)", F_TITLE, align=LEFT, border=False)
    cell(ws, "A2", "a_ij = цена_j / цена_i, ограничено шкалой Саати [1/9; 9]", align=LEFT, border=False)
    cell(ws, "A4", "Поставщик", F_BOLD, GREY)
    cell(ws, "B4", "Цена, ₽", F_BOLD, GREY)
    for j in range(m):
        cell(ws, f"{get_column_letter(3 + j)}4", ups.loc[j, "vendor"], F_BOLD, GREY)
    gcol, wcol, awcol, rcol = (get_column_letter(3 + m + k) for k in range(4))
    for ref, t in ((gcol, "Ср. геом."), (wcol, "Приоритет"), (awcol, "(A·w)"), (rcol, "(A·w)/w")):
        cell(ws, f"{ref}4", t, F_BOLD, GREY)
    first, last = 5, 4 + m
    for i in range(m):
        r = first + i
        cell(ws, f"A{r}", ups.loc[i, "vendor"], F_BOLD, align=LEFT)
        cell(ws, f"B{r}", f"=Данные!D{data_rows[i]}", F_LINK, fmt="#,##0")
        for j in range(m):
            cell(ws, f"{get_column_letter(3 + j)}{r}", f"=MIN(9,MAX(1/9,$B${first + j}/$B{r}))", fmt="0.00")
        rng = f"{get_column_letter(3)}{r}:{get_column_letter(2 + m)}{r}"
        cell(ws, f"{gcol}{r}", f"=PRODUCT({rng})^(1/{m})", fmt="0.0000")
        cell(ws, f"{wcol}{r}", f"={gcol}{r}/SUM(${gcol}${first}:${gcol}${last})", F_BOLD, fmt="0.000")
        terms = "+".join(f"{get_column_letter(3 + j)}{r}*${wcol}${first + j}" for j in range(m))
        cell(ws, f"{awcol}{r}", f"={terms}", fmt="0.0000")
        cell(ws, f"{rcol}{r}", f"={awcol}{r}/{wcol}{r}", fmt="0.0000")
    s = last + 2
    ri = {3: 0.58, 4: 0.90, 5: 1.12}.get(m, 0.9)
    for k, (label, value, fmt) in enumerate([
            ("λmax", f"=AVERAGE({rcol}{first}:{rcol}{last})", "0.0000"),
            ("ИС", f"=(B{s}-{m})/({m}-1)", "0.0000"),
            ("СИ", ri, "0.00"),
            ("ОС", f"=IF(B{s + 2}=0,0,B{s + 1}/B{s + 2})", "0.0000")]):
        cell(ws, f"A{s + k}", label, F_BOLD, GREY, align=LEFT)
        cell(ws, f"B{s + k}", value, F_INPUT if k == 2 else F_TEXT, fmt=fmt)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 13
    for k in range(m + 4):
        ws.column_dimensions[get_column_letter(3 + k)].width = 11


def main():
    cfg = json.loads((ROOT / "config" / "ahp_criteria.json").read_text(encoding="utf-8"))
    criteria = cfg["criteria"]
    ids = [c["id"] for c in criteria]
    judgments = {(ids.index(a), ids.index(b)): v for a, b, v in cfg["judgments"]}
    py_weights = pd.read_csv(OUT / "criteria_weights.csv")["weight"].tolist()
    scored = pd.read_csv(ROOT / "data" / "processed" / "alternatives_scored.csv", dtype=str).fillna("")
    local = pd.read_csv(OUT / "local_priorities.csv")

    wb = Workbook()
    sheet_instructions(wb)
    info = sheet_criteria(wb, criteria, judgments, py_weights)
    sheet_data(wb, scored)
    sheet_priorities(wb, local, ids, info["weight_cells"])
    sheet_example(wb, scored)
    path = OUT / "ahp_model.xlsx"
    wb.save(path)
    print("saved", path)


if __name__ == "__main__":
    main()
